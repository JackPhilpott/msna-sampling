"""Step 5: one Excel workbook per partner.

Each partner gets what its field teams need: a summary of its sample, the
clusters to visit, the household list (6 primary and 6 reserve per cluster)
and the resampling history. Every sheet carries the synthetic-data label.

Inputs: data/synthetic/{clusters_after_resampling,resample_log}.csv,
        data/synthetic/households/*.csv
Outputs: data/synthetic/partner_workbooks/<partner>.xlsx
"""
import re
import sys
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python"))
from sampling_tools import SYNTHETIC_LABEL, load_config, output_dir  # noqa: E402

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LABEL_FONT = Font(bold=True, color="C00000", size=12)


def write_sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame, title: str) -> None:
    """Write a table under a two-line heading: the sheet title and the synthetic label."""
    df.to_excel(writer, sheet_name=name, index=False, startrow=3)
    ws = writer.sheets[name]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = SYNTHETIC_LABEL
    ws["A2"].font = LABEL_FONT
    for cell in ws[4]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.freeze_panes = "A5"
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(45, width + 2)


def make_reproducible(path: Path) -> None:
    """Fix the timestamps an .xlsx file carries, so the same content gives the same file."""
    with zipfile.ZipFile(path) as z:
        parts = [(info, z.read(info.filename)) for info in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for info, data in parts:
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            z.writestr(fixed, data)


def main() -> int:
    cfg = load_config()
    out = output_dir(cfg)
    clusters = pd.read_csv(out / "clusters_after_resampling.csv")
    log = pd.read_csv(out / "resample_log.csv")
    households = pd.concat([pd.read_csv(f) for f in sorted((out / "households").glob("*.csv"))],
                           ignore_index=True)
    dest = out / "partner_workbooks"
    dest.mkdir(exist_ok=True)
    for old in dest.glob("*.xlsx"):
        old.unlink()

    for partner, cl in clusters.groupby("partner", sort=True):
        active = cl[cl["status"] == "active"]
        hh = households[households["partner"] == partner]
        lg = log[log["partner"] == partner]
        summary = pd.DataFrame({
            "Item": ["Partner", "States", "LGAs", "Active clusters", "  of which replacements",
                     "Households to interview (primary)", "Reserve households",
                     "Clusters reported inaccessible", "Clusters that could not be replaced (shortfall)"],
            "Value": [partner, ", ".join(sorted(cl["adm1_name"].unique())), cl["adm2_name"].nunique(),
                      len(active), int((active["round_added"] > 0).sum()),
                      int((hh["role"] == "primary").sum()), int((hh["role"] == "reserve").sum()),
                      int((cl["status"] == "inaccessible").sum()), int(lg["shortfall"].sum())],
        })
        visit = active[["cluster_id", "stratum_id", "pop_type", "adm1_name", "adm2_name", "psu_id",
                        "draw", "lon", "lat"]].sort_values("cluster_id")
        hh_list = hh[["household_id", "cluster_id", "role", "rank", "adm2_name", "lon", "lat"]] \
            .sort_values(["cluster_id", "role", "rank"], ascending=[True, False, True])
        history = cl[cl["status"] == "inaccessible"][["cluster_id", "stratum_id", "adm2_name", "round_removed"]] \
            .merge(cl[cl["replaces"].notna() & (cl["replaces"] != "")][["replaces", "cluster_id"]]
                   .rename(columns={"replaces": "cluster_id", "cluster_id": "replaced_by"}),
                   on="cluster_id", how="left").fillna({"replaced_by": "not replaced (shortfall)"})

        path = dest / f"{partner.lower().replace(' ', '_')}.xlsx"
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            write_sheet(writer, "Summary", summary, f"MSNA sample: {partner}")
            write_sheet(writer, "Clusters", visit, "Clusters to visit")
            write_sheet(writer, "Households", hh_list, "Households: 6 primary and 6 reserve per cluster")
            write_sheet(writer, "Resampling", history, "Clusters reported inaccessible and their replacements")
        make_reproducible(path)

    print(f"Partner workbooks: {clusters['partner'].nunique()} written to {cfg['outputs']['dir']}/partner_workbooks")
    return 0


if __name__ == "__main__":
    sys.exit(main())
